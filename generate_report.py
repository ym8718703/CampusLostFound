"""
Simple Report Generator for Campus Lost & Found System
Uses standard python-docx library to create the project report document (.docx).
Clean, readable, and easy to understand without any complex XML manipulation.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_report():
    # 1. Initialize Document
    doc = docx.Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # Helper Functions for Clean Document Formatting
    # -------------------------------------------------------------
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 54, 93) # Deep Blue

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.italic = True
        run.font.color.rgb = RGBColor(100, 100, 100)

    def add_heading1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 54, 93)

    def add_heading2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(43, 108, 176)

    def add_heading3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.bold = True

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)

    def add_table(headers, rows):
        """Creates a clean, styled table using built-in Word Table Grid"""
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.style = 'Table Grid'
        table.autofit = True

        # Header row
        for col_idx, header_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = header_text
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
                    run.font.bold = True

        # Data rows
        for row_idx, row_data in enumerate(rows):
            for col_idx, cell_value in enumerate(row_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(cell_value)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_code(code_text):
        """Adds code snippet in monospaced font"""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 30, 30)

    # -------------------------------------------------------------
    # DOCUMENT CONTENT
    # -------------------------------------------------------------

    # College Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("JSPM's Rajarshi Shahu College of Engineering, Tathawade, Pune - 411033\n"
                            "(An Autonomous Institute Affiliated to Savitribai Phule Pune University)\n"
                            "Department of Computer Science and Business Systems\n"
                            "CB2206L: Skill Development Lab-I (Pattern 2025-26) | Semester III\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(43, 108, 176)

    add_title("Mini Project Report\nCampus Lost & Found Management System")
    add_subtitle("A Syllabus-Aligned Core Java Application (CO2: OOP | CO3: Strings & Exceptions | CO4: Collections & Multithreading)")

    # Group Members Table
    add_table(["Sr. No.", "Student Name", "PRN Number"], [
        ["1.", "[Student Name 1]", "[PRN No. 1]"],
        ["2.", "[Student Name 2]", "[PRN No. 2]"],
        ["3.", "[Student Name 3]", "[PRN No. 3]"],
        ["4.", "[Student Name 4]", "[PRN No. 4]"]
    ])

    # Section 1
    add_heading1("1. Title of the Project")
    add_p("Campus Lost & Found Management System", bold_prefix="Project Title: ")
    add_p("A modular Core Java application providing both an interactive menu-driven Console interface and a Java Swing Graphical User Interface (GUI) to report lost and found items on campus, persist records in text files, search items using overloaded methods, execute concurrent background matching with thread synchronization, and match found items with reported lost items using simple Java concepts.")

    # Section 2
    add_heading1("2. Problem Statement")
    add_p("On a university campus, students and staff frequently misplace personal belongings such as student ID cards, wallets, calculators, notebooks, keys, clothing, and water bottles. Traditionally, lost and found reports are handled informally through handwritten register notices, WhatsApp messages, or verbal enquiries at the security desk. This informal workflow exhibits several drawbacks:")
    add_bullet("No single, organized repository exists to track what has been lost and found.", bold_prefix="Lack of Centralization: ")
    add_bullet("Important details such as location of loss/discovery, timestamp, item description, and contact info are often incomplete.", bold_prefix="Incomplete Information: ")
    add_bullet("Matching relies entirely on manual comparison or memory, leaving many items unclaimed.", bold_prefix="Manual Overhead: ")
    add_bullet("No formal status tracking exists to confirm whether an item was claimed and returned.", bold_prefix="No Claim History: ")
    add_p("To solve this problem, a simple and robust Core Java application was developed to allow structured reporting, persistent file storage, input validation, and automatic matching between lost and found items.")

    # Section 3
    add_heading1("3. Objectives")
    add_p("The project is strictly designed around the CB2206L Course Outcomes (COs):")
    add_bullet("Apply Object-Oriented Programming (CO2) using classes, objects, parameterized constructors, inheritance, user-defined packages, access modifiers, abstraction (abstract classes & interfaces), method overloading, and method overriding.", bold_prefix="CO2 (OOP Concepts): ")
    add_bullet("Implement String Processing and Exception Handling (CO3) using standard Java String methods (validating contact and dates) and user-defined checked exceptions (InvalidItemException, ItemNotFoundException).", bold_prefix="CO3 (Strings & Exceptions): ")
    add_bullet("Utilize the Java Collections Framework (CO4) using ArrayList for storing items, HashMap for fast O(1) item ID retrieval, and HashSet for collecting unique categories.", bold_prefix="CO4 (Collections Framework): ")
    add_bullet("Implement Multithreading and Thread Synchronization (CO4, Exp 5) by developing a background worker thread (AutoMatchThread extending Thread) that safely synchronizes on shared inventory to scan and export match reports.", bold_prefix="CO4 (Multithreading & Sync): ")
    add_bullet("Provide a dual-interface system: a primary menu-driven Console interface (via Scanner) for standard lab viva and an optional Swing desktop GUI.", bold_prefix="User Interface: ")
    add_bullet("Ensure File Persistence using standard Core Java I/O (BufferedReader, BufferedWriter, PrintWriter) without external databases.", bold_prefix="File Handling: ")
    add_bullet("Maintain strict syllabus alignment by replacing complex or out-of-syllabus features (LocalDate, ChronoUnit, Lambdas, Streams, complex regex) with simple, standard Core Java constructs.", bold_prefix="Syllabus Strictness: ")

    # Section 4
    add_heading1("4. Methodology")
    add_heading2("4.1 Workflow of the System")
    add_p("When the system starts, FileManager reads data/lost_items.txt and data/found_items.txt, restoring records into memory. The user interacts through the console menu (Main.java) or the Swing GUI (MainFrame.java):")
    add_bullet("Step 1 (Load Records): The system loads saved records into an ArrayList<LostItem>, ArrayList<FoundItem>, and a HashMap<String, Item> index.", bold_prefix="1. Initialization: ")
    add_bullet("Step 2 (Report Item & Validate): The user inputs item details. Validated setters check the inputs. If invalid, an InvalidItemException is thrown and caught with a clear error message. If valid, an ID (e.g. L001, F001) is assigned.", bold_prefix="2. Input Validation: ")
    add_bullet("Step 3 (Immediate Match Alert): The system checks if the new item matches any existing opposite item and alerts the user.", bold_prefix="3. Match Check: ")
    add_bullet("Step 4 (Overloaded Search): Users can search items using two overloaded methods: search(keyword) or search(keyword, category).", bold_prefix="4. Search: ")
    add_bullet("Step 5 (Claim Confirmation): When an item is returned, confirmMatch() sets the status of both lost and found records to RETURNED and saves to file.", bold_prefix="5. Return Item: ")
    add_bullet("Step 6 (Multithreaded Background Sync): Users can start AutoMatchThread (extends Thread), which synchronizes on ItemManager, computes matches, and exports data/match_report.txt in the background.", bold_prefix="6. Multithreading: ")

    add_heading2("4.2 Matching Algorithm Using Simple Java Concepts")
    add_p("Matching between lost items and found items is implemented in LostItem.java using simple, standard Java String concepts (equalsIgnoreCase, toLowerCase, split, and contains):")
    add_table(["Criterion", "Simple Java Implementation", "Points Awarded"], [
        ["Category Match", "Compares categories using String.equalsIgnoreCase()", "40 Points"],
        ["Name / Description Word Match", "Splits the item name by space using split(' '). Checks if any meaningful word (length >= 3) appears in the other item using String.contains().", "30 Points"],
        ["Location Match", "Checks if locations match using equalsIgnoreCase() or contains(). Exact match gives 30 points; partial match gives 20 points.", "Up to 30 Points"],
        ["Safety Rule", "If no name or description word matches, the score is 0. Total score is 0 to 100. Any pair with score >= 50% is reported as a match.", "Threshold: 50%"]
    ])

    # Section 5
    add_heading1("5. Dataset, Tools & Technologies Used")
    add_heading2("5.1 Dataset & File Storage")
    add_p("The project uses plain text files stored inside the data/ folder:")
    add_bullet("data/lost_items.txt: Stores all lost item records.", bold_prefix="Lost Items: ")
    add_bullet("data/found_items.txt: Stores all found item records.", bold_prefix="Found Items: ")
    add_bullet("data/match_report.txt: Stores the exported match report.", bold_prefix="Match Report: ")
    add_p("Each line contains nine fields separated by the pipe character (|):\n"
          "Example: L001|Black Wallet|Wallet|Leather wallet with college ID card|Library|2026-10-10|Aarav Patil|9876543210|OPEN")

    add_table(["Field Index & Name", "Description", "Simple Validation Rule"], [
        ["f[0] : id", "Unique ID (String)", "Auto-generated: L001, L002... or F001, F002..."],
        ["f[1] : name", "Item Name (String)", "Non-empty string; validated using trim().length() > 0."],
        ["f[2] : category", "Category (String)", "Electronics, Wallet, Books, ID Card, Keys, Clothing, Other."],
        ["f[3] : description", "Description (String)", "Optional details (color, brand, marks)."],
        ["f[4] : location", "Location (String)", "Non-empty string."],
        ["f[5] : date", "Date (String: YYYY-MM-DD)", "Validated by splitting with '-' and Integer.parseInt()."],
        ["f[6] : personName", "Owner / Finder Name (String)", "Non-empty string."],
        ["f[7] : personContact", "Contact Info (String)", "10-digit number (all digits) or email containing '@' and '.'."],
        ["f[8] : status", "Status (String)", "Must be 'OPEN' or 'RETURNED'."]
    ])

    add_heading2("5.2 Tools & Technologies")
    add_table(["Component", "Technology / Tool Used"], [
        ["Programming Language", "Java Standard Edition (JDK 8 or later)"],
        ["Package", "User-defined package: campuslostfound (Satisfies CO2)"],
        ["User Interface", "1. Menu-Driven Console Application (java.util.Scanner)\n2. Desktop GUI (Java Swing with classic ActionListener inner classes)"],
        ["Collections", "ArrayList (inventory lists), HashMap (fast ID lookup), HashSet (unique categories)"],
        ["Multithreading", "Thread class, Runnable, synchronized block, Thread.sleep()"],
        ["File Handling", "BufferedReader, BufferedWriter, FileReader, FileWriter, PrintWriter"],
        ["Exception Handling", "Custom checked exceptions: InvalidItemException, ItemNotFoundException"]
    ])

    add_heading2("5.3 Out-of-Syllabus Concepts Replaced with Simple Java")
    add_table(["Out-of-Syllabus Concept", "Reason for Removal", "Simple Syllabus-Compliant Replacement"], [
        ["java.time.LocalDate, ChronoUnit", "Modern Java 8 Date API is not in Core Java syllabus.", "Standard String format (YYYY-MM-DD) validated using split('-') and Integer.parseInt()."],
        ["Java 8 Lambda Expressions (() -> ...)", "Functional programming is outside syllabus scope.", "Standard Scanner menu flow and classic ActionListener inner classes."],
        ["Java 8 Stream API & Method References", "Streams are outside the Core Java syllabus.", "Standard for-loops and Collections.sort() with Comparable."],
        ["Complex Regex Patterns", "Regex engines are not part of Unit II String processing.", "Standard String methods: length(), charAt(), Character.isDigit(), indexOf(), split(' ')."],
        ["Complex NLP Stop-Words & Token Parsing", "Too complicated for 2nd-year Core Java.", "Simple String split(' ') and contains() for keyword matching."]
    ])

    # Section 6
    add_heading1("6. Implementation")
    add_heading2("6.1 Class Design and Responsibilities")
    add_table(["Class / Interface", "Type", "Responsibility & Syllabus Mapping"], [
        ["Item", "Abstract Class", "Base class with encapsulated fields, abstract methods, overloaded displayDetails(), and validation (CO2, CO3)."],
        ["LostItem", "Subclass (extends Item)", "Represents lost items; implements Matchable with simple String matching (CO2, CO3)."],
        ["FoundItem", "Subclass (extends Item)", "Represents found items; overrides abstract methods (CO2)."],
        ["Storable", "Interface", "Defines toFileString() for file persistence (CO2)."],
        ["Matchable", "Interface", "Defines MATCH_THRESHOLD = 50 and calculateMatchScore(Item) (CO2)."],
        ["MatchResult", "Comparable Class", "Holds a matched pair and score; implements Comparable for sorting (CO4)."],
        ["ItemManager", "Controller Class", "Manages ArrayList, HashMap, and HashSet; provides overloaded search() and synchronized methods (CO2, CO4)."],
        ["AutoMatchThread", "Thread Class", "Extends Thread; implements run() with synchronized block to scan and export reports in background (CO4, Exp 5)."],
        ["FileManager", "File Handling Class", "Reads and writes records using BufferedReader, BufferedWriter, and PrintWriter."],
        ["InvalidItemException", "Custom Exception", "User-defined checked exception thrown when user input is invalid (CO3)."],
        ["ItemNotFoundException", "Custom Exception", "User-defined checked exception thrown when an item ID does not exist (CO3)."],
        ["Main", "Console UI Class", "Entry point providing an interactive 10-option menu-driven console interface (CO1-CO4)."],
        ["MainFrame", "Swing GUI Class", "Desktop GUI window with tabs for reporting, viewing, and matching."]
    ])

    add_heading2("6.2 Key Code Excerpt: Simplified Java Matching (LostItem.java)")
    add_p("The matching method uses only standard String methods (equalsIgnoreCase, toLowerCase, split, and contains):")
    add_code("""@Override
public int calculateMatchScore(Item other) {
    if (!(other instanceof FoundItem)) {
        return 0;
    }

    int score = 0;

    // 1. Category Match (40 points) using equalsIgnoreCase()
    boolean sameCategory = getCategory().equalsIgnoreCase(other.getCategory());
    if (sameCategory) {
        score += 40;
    }

    // 2. Name Word Match (30 points) using split(" ") and contains()
    String myName = getName().toLowerCase();
    String otherText = (other.getName() + " " + other.getDescription()).toLowerCase();

    boolean nameMatched = false;
    if (otherText.contains(myName) || myName.contains(other.getName().toLowerCase())) {
        nameMatched = true;
    } else {
        String[] words = myName.split(" ");
        for (int i = 0; i < words.length; i++) {
            String word = words[i].trim();
            if (word.length() >= 3 && otherText.contains(word)) {
                nameMatched = true;
                break;
            }
        }
    }

    if (nameMatched) {
        score += 30;
    }

    // 3. Location Match (up to 30 points)
    String myLoc = getLocation().toLowerCase();
    String otherLoc = other.getLocation().toLowerCase();
    if (myLoc.equalsIgnoreCase(otherLoc)) {
        score += 30; // Exact match
    } else if (myLoc.contains(otherLoc) || otherLoc.contains(myLoc)) {
        score += 20; // Partial match
    }

    // Safety rule: Must share at least one name word
    if (!nameMatched) {
        return 0;
    }

    return Math.min(score, 100);
}""")

    add_heading2("6.3 Key Code Excerpt: Multithreading & Synchronization (AutoMatchThread.java)")
    add_p("Extends Thread, overrides run(), and uses a synchronized block on ItemManager:")
    add_code("""public class AutoMatchThread extends Thread {
    private final ItemManager manager;
    private final FileManager fileManager;

    public AutoMatchThread(ItemManager manager, FileManager fileManager) {
        super("AutoMatchWorkerThread");
        this.manager = manager;
        this.fileManager = fileManager;
    }

    @Override
    public void run() {
        System.out.println("\\n[" + getName() + "] Background thread started.");
        try {
            System.out.println("[" + getName() + "] Step 1: Scanning inventory...");
            Thread.sleep(600);

            List<MatchResult> matches;
            // Thread synchronization on shared ItemManager instance
            synchronized (manager) {
                System.out.println("[" + getName() + "] Step 2: Acquired lock. Computing matches...");
                matches = manager.findAllMatches();
            }

            Thread.sleep(400);
            System.out.println("[" + getName() + "] Step 3: Found " + matches.size() + " match(es).");
            fileManager.exportMatchReport(matches, FileManager.REPORT_FILE);
            System.out.println("[" + getName() + "] Step 4: Saved report to: " + FileManager.REPORT_FILE);
        } catch (InterruptedException | IOException e) {
            System.err.println("[" + getName() + "] Error: " + e.getMessage());
        }
    }
}""")

    add_heading2("6.4 How to Compile and Run the Project")
    add_code("""# 1. Compile all Java files from project root:
javac campuslostfound/*.java

# 2. Run the interactive Console Menu application:
java campuslostfound.Main

# 3. (Optional) Run the Swing GUI directly:
java campuslostfound.Main --gui

# 4. Or double-click 'run.bat' on Windows to compile and run automatically.""")

    # Section 7
    add_heading1("7. Results & Test Verification")
    add_heading2("7.1 Matching Results on Sample Data")
    add_table(["Lost Item", "Found Item", "Score", "Matching Rationale"], [
        ["L003: Casio Calculator (Exam Hall)", "F002: Calculator (Exam Hall)", "100%", "Category match (40), word 'calculator' match (30), exact location match (30)."],
        ["L004: House Keys (Parking Area)", "F003: Keys (Parking Area)", "100%", "Category match (40), word 'keys' match (30), exact location match (30)."],
        ["L006: Data Structures Book (Computer Lab)", "F005: Textbook (Computer Lab)", "100%", "Category match (40), word 'book' match in description (30), location match (30)."],
        ["L001: Black Wallet (Library)", "F001: Wallet (Library 2nd Floor)", "90%", "Category match (40), word 'wallet' match (30), partial location match (20)."],
        ["L002: Blue Water Bottle (Canteen)", "F004: Red Umbrella (Canteen)", "0% (No match)", "Zero words in common; correctly rejected by the safety rule."]
    ])

    add_heading2("7.2 Test Cases Suite (14 Comprehensive Tests)")
    add_table(["TC #", "Test Case Description", "Input / Action", "Expected Result", "Status"], [
        ["TC-01", "Register Lost Item", "Casio Calc, Electronics, 2026-10-15, 9876543210", "Item assigned ID L007, saved to file.", "Pass"],
        ["TC-02", "Immediate Match Alert", "Submit report matching existing found item", "System immediately alerts user of potential match.", "Pass"],
        ["TC-03", "Empty Name Validation", "Submit empty item name", "InvalidItemException: 'Item name cannot be empty.'", "Pass"],
        ["TC-04", "Invalid Date Validation", "Enter date '15/10/2026'", "InvalidItemException: 'Date must be in YYYY-MM-DD format.'", "Pass"],
        ["TC-05", "Invalid Contact Validation", "Enter phone '12345'", "InvalidItemException: 'Contact must be 10 digits or email.'", "Pass"],
        ["TC-06", "Overloaded Search (Keyword)", "search('wallet')", "Returns items containing 'wallet' in details.", "Pass"],
        ["TC-07", "Overloaded Search (Cat Filter)", "search('Keys', 'Keys')", "Returns items matching keyword and category filter.", "Pass"],
        ["TC-08", "HashMap ID Lookup", "getItemById('L001')", "Returns item L001 in O(1) time.", "Pass"],
        ["TC-09", "HashSet Categories Listing", "Select Menu Option 8", "Displays set of unique categories without duplicates.", "Pass"],
        ["TC-10", "Multithreaded Sync", "Select Menu Option 9", "AutoMatchThread runs concurrently and exports report.", "Pass"],
        ["TC-11", "Confirm Return", "Confirm match L003 <-> F002", "Both items marked RETURNED and saved to file.", "Pass"],
        ["TC-12", "Delete Record", "Delete item 'L005'", "Item removed from ArrayList, HashMap, and file.", "Pass"],
        ["TC-13", "File Persistence", "Restart application", "All records successfully reloaded from text files.", "Pass"],
        ["TC-14", "Missing ID Exception", "Delete unknown ID 'L999'", "ItemNotFoundException: 'Item L999 was not found.'", "Pass"]
    ])

    # Section 8
    add_heading1("8. Conclusion")
    add_p("The Campus Lost & Found Management System successfully satisfies all requirements of the CB2206L syllabus:")
    add_bullet("Fully covers OOP (CO2), String Processing & Exceptions (CO3), Collections & Multithreading (CO4).", bold_prefix="1. Syllabus Compliance: ")
    add_bullet("Uses simple, understandable Java concepts throughout, making the project easy to explain in lab exams and viva.", bold_prefix="2. Simplicity & Clarity: ")
    add_bullet("Provides dual interfaces: a transparent Console Menu for evaluation and a Swing GUI for visual interaction.", bold_prefix="3. Usability: ")
    add_bullet("Guarantees data persistence through text file handling with zero third-party dependencies.", bold_prefix="4. Independence: ")

    # Appendix: Source Code
    add_heading1("Appendix: Complete Source Code")
    add_p("The complete source code for all 13 Java classes under package campuslostfound is given below:")

    code_files = [
        ("Item.java", "campuslostfound/Item.java"),
        ("LostItem.java", "campuslostfound/LostItem.java"),
        ("FoundItem.java", "campuslostfound/FoundItem.java"),
        ("Storable.java", "campuslostfound/Storable.java"),
        ("Matchable.java", "campuslostfound/Matchable.java"),
        ("MatchResult.java", "campuslostfound/MatchResult.java"),
        ("InvalidItemException.java", "campuslostfound/InvalidItemException.java"),
        ("ItemNotFoundException.java", "campuslostfound/ItemNotFoundException.java"),
        ("FileManager.java", "campuslostfound/FileManager.java"),
        ("ItemManager.java", "campuslostfound/ItemManager.java"),
        ("AutoMatchThread.java", "campuslostfound/AutoMatchThread.java"),
        ("Main.java", "campuslostfound/Main.java"),
        ("MainFrame.java", "campuslostfound/MainFrame.java")
    ]

    base_dir = r"C:\Users\yashr\OneDrive\Documents\Academic\SY\SEM-1\SDL\CampusLostFound"
    for fname, rel_path in code_files:
        full_path = os.path.join(base_dir, rel_path)
        if os.path.exists(full_path):
            with open(full_path, "r", encoding="utf-8") as f:
                code_text = f.read()
            add_heading2(f"Source Code: {fname} (package campuslostfound)")
            add_code(code_text)

    # Save Word document
    output_docx = os.path.join(base_dir, "Campus_Lost_Found_Mini_Project_Report.docx")
    doc.save(output_docx)
    print(f"[SUCCESS] Word report created at: {output_docx}")

if __name__ == "__main__":
    build_report()
